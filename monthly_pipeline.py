#!/usr/bin/env python3
"""
Orchestrator for the scheduled monthly workflow.

LTA publishes each month's data by the 10th of the following month. This
script is meant to run daily from the 11th onward: it first checks whether
the just-published month is already present in the committed workbook
(no wasted API calls if so), and only calls the LTA API when there's
actually new data to fetch. Sets GITHUB_OUTPUT `updated=true|false` for the
workflow to decide whether to commit, and exits non-zero when the new
month isn't published yet so the workflow's daily retry has something to
act on.
"""

import csv
import os
import sys
from datetime import date

from fetch_lta_data import fetch_all_historical
from build_reports import build_workbook

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

FULL_XLSX = "lta_train_od_historical_by_month.xlsx"
EW1_XLSX = "lta_train_od_EW1_by_month.xlsx"
CSV_TMP = "lta_train_od_historical.csv"
STATION = "EW1"


def target_month():
    """The most recently completed calendar month, as (YYYYMM, YYYY-MM)."""
    today = date.today()
    year, month = today.year, today.month
    month -= 1
    if month == 0:
        month = 12
        year -= 1
    return f"{year}{month:02d}", f"{year}-{month:02d}"


def already_have_month(dash_month):
    if load_workbook is None or not os.path.exists(FULL_XLSX):
        return False
    wb = load_workbook(FULL_XLSX, read_only=True)
    return dash_month in wb.sheetnames


def months_in_csv(csv_path):
    with open(csv_path, newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        header = next(reader)
        idx = header.index("YEAR_MONTH")
        return {row[idx] for row in reader}


def set_output(name, value):
    path = os.environ.get("GITHUB_OUTPUT")
    if path:
        with open(path, "a") as f:
            f.write(f"{name}={value}\n")


def main():
    api_key = sys.argv[1] if len(sys.argv) > 1 else None
    if not api_key:
        print("Usage: python3 monthly_pipeline.py <api_key>")
        sys.exit(2)

    yyyymm, dash_month = target_month()

    if already_have_month(dash_month):
        print(f"Already have data for {dash_month} — nothing to do.")
        set_output("updated", "false")
        sys.exit(0)

    print(f"Looking for newly published data for {dash_month}...")
    success = fetch_all_historical(api_key, CSV_TMP, months_back=4)

    if not success:
        print(f"No data available yet for {dash_month}.")
        set_output("updated", "false")
        sys.exit(1)

    found_months = months_in_csv(CSV_TMP)
    if dash_month not in found_months:
        print(f"Fetch succeeded but {dash_month} still isn't published (found: {sorted(found_months)}).")
        set_output("updated", "false")
        sys.exit(1)

    build_workbook(CSV_TMP, FULL_XLSX)
    build_workbook(CSV_TMP, EW1_XLSX, station=STATION)
    print(f"Updated with {dash_month} data.")
    set_output("updated", "true")


if __name__ == "__main__":
    main()
